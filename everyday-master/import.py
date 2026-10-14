import os
import openpyxl
import psycopg2
from config import DB_CONFIG, DATA_DIR

def import_sizes(cur):
    file_path = os.path.join(DATA_DIR, "Sizes_import.xlsx")
    if not os.path.exists(file_path):
        print(f"⚠️ Файл не найден: {file_path}")
        return

    wb = openpyxl.load_workbook(file_path)
    sheet = wb.active
   
    count = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            size_val = float(row[0])
            cur.execute(
                "INSERT INTO sizes (value) VALUES (%s) ON CONFLICT (value) DO NOTHING;",
                (size_val,)
            )
            count += 1
    print(f"✅ Импортировано размеров: {count}")

def import_users(cur):
    file_path = os.path.join(DATA_DIR, "Users_import.xlsx")
    if not os.path.exists(file_path):
        print(f"⚠️ Файл не найден: {file_path}")
        return

    wb = openpyxl.load_workbook(file_path)
    sheet = wb.active
   
    count = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[3] is not None:  # Логин
            last_name, first_name, middle_name, login, role = row[0], row[1], row[2], row[3], row[4]
            cur.execute(
                """
                INSERT INTO users (last_name, first_name, middle_name, login, role)
                VALUES (%s, %s, %s, %s, %s)
                ON CONFLICT (login) DO NOTHING;
                """,
                (last_name, first_name, middle_name, login, role)
            )
            count += 1
    print(f"✅ Импортировано пользователей: {count}")

def import_products(cur):
    file_path = os.path.join(DATA_DIR, "Products_import.xlsx")
    if not os.path.exists(file_path):
        print(f"⚠️ Файл не найден: {file_path}")
        return

    wb = openpyxl.load_workbook(file_path)
    sheet = wb.active
   
    count = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[3] is not None:  # Наименование товара
            category, subcategory, image, product_name, manufacturer, description, composition, price = row
            cur.execute(
                """
                INSERT INTO products (category, subcategory, image, product_name, manufacturer, description, composition, price)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s);
                """,
                (category, subcategory, image, product_name, manufacturer, description, composition, price)
            )
            count += 1
    print(f"✅ Импортировано товаров: {count}")

def import_stock(cur):
    file_path = os.path.join(DATA_DIR, "Stock_Items_import.xlsx")
    if not os.path.exists(file_path):
        print(f"⚠️ Файл не найден: {file_path}")
        return

    wb = openpyxl.load_workbook(file_path)
    sheet = wb.active
   
    count = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            product_name, manufacturer, size, stock_quantity = row
            cur.execute(
                """
                INSERT INTO stock_items (product_name, manufacturer, size, stock_quantity)
                VALUES (%s, %s, %s, %s);
                """,
                (product_name, manufacturer, size, stock_quantity)
            )
            count += 1
    print(f"✅ Импортировано складских позиций: {count}")

def import_orders(cur):
    file_path = os.path.join(DATA_DIR, "Orders_import.xlsx")
    if not os.path.exists(file_path):
        print(f"⚠️ Файл не найден: {file_path}")
        return

    wb = openpyxl.load_workbook(file_path)
    sheet = wb.active
   
    count = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            order_number, order_date, customer_name, category, product_name, manufacturer, size, quantity, unit_price = row
            cur.execute(
                """
                INSERT INTO orders (order_number, order_date, customer_name, category, product_name, manufacturer, size, quantity, unit_price)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s);
                """,
                (order_number, order_date, customer_name, category, product_name, manufacturer, size, quantity, unit_price)
            )
            count += 1
    print(f"✅ Импортировано заказов: {count}")

def main():
    print("==================================================")
    print("Импорт данных в БД Чудо Обувь")
    print("==================================================")
   
    conn = psycopg2.connect(**DB_CONFIG)
    conn.autocommit = False
    cur = conn.cursor()
   
    try:
        print("\n[1/5] Импорт размеров...")
        import_sizes(cur)
       
        print("\n[2/5] Импорт пользователей...")
        import_users(cur)
       
        print("\n[3/5] Импорт товаров...")
        import_products(cur)
       
        print("\n[4/5] Импорт складских остатков...")
        import_stock(cur)
       
        print("\n[5/5] Импорт заказов...")
        import_orders(cur)
       
        conn.commit()
        print("\n🎉 Все данные успешно импортированы в базу данных!")
    except Exception as e:
        conn.rollback()
        print(f"\n❌ Произошла ошибка при импорте: {e}")
    finally:
        cur.close()
        conn.close()

if __name__ == "__main__":
    main()